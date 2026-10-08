"""Generate readable Word guides from the maintained phase-2 Markdown notes."""
from pathlib import Path
import re
import sys

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / "docs" / "phase2"


def plain(value):
    value = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", value)
    return value.replace("**", "").replace("`", "")


def markdown_document(source, target):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[name]
        style.font.name = "Noto Sans CJK SC"
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    normal = document.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(5)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        document.styles[name].font.color.rgb = RGBColor.from_string("17365D")
    section.header.paragraphs[0].text = "电动工具预测性维护研发 · 第二阶段软件与公开数据预验证"
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("2026-10-09  |  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("```"):
            language = line[3:]
            code = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                code.append(lines[index])
                index += 1
            if language == "mermaid":
                document.add_paragraph("公开波形 → 重采样／分区／切窗 → 去均值／小波处理 → 轻量 CNN → 可选注意力 → 四类故障部位。FFT 特征另保存为解释性特征表。")
            else:
                paragraph = document.add_paragraph("\n".join(code))
                paragraph.paragraph_format.line_spacing = 1.1
                for run in paragraph.runs:
                    run.font.name, run.font.size = "Consolas", Pt(8)
            index += 1
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [plain(value.strip()) for value in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[:\- ]+", value) for value in cells):
                    rows.append(cells)
                index += 1
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Light Shading Accent 1"
            for i, row in enumerate(rows):
                for j, cell in enumerate(row):
                    table.cell(i, j).text = cell
                    for paragraph in table.cell(i, j).paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(8)
            document.add_paragraph()
            continue
        if line.startswith("# "):
            document.add_heading(plain(line[2:]), 0)
        elif line.startswith("### "):
            document.add_heading(plain(line[4:]), 2)
        elif line.startswith("## "):
            document.add_heading(plain(line[3:]), 1)
        elif re.match(r"^\d+\. ", line):
            document.add_paragraph(plain(line), style="Normal")
        elif line.startswith("- "):
            document.add_paragraph(plain(line[2:]), style="List Bullet")
        else:
            document.add_paragraph(plain(line))
        index += 1
    if source.name == "早上先看这里.md":
        for image_name, caption in (("公开轴承分析.png", "公开数据：状态、算法判断与真实标签分别展示"),
                                    ("虚拟实验台.png", "虚拟台：生成条件与三个传感器分开展示")):
            path = FOLDER / "网页截图" / image_name
            if path.exists():
                document.add_page_break()
                document.add_heading(caption, 1)
                document.add_picture(str(path), width=Cm(17))
    document.core_properties.title = plain(lines[0].lstrip("# "))
    document.core_properties.subject = "真实公开数据预验证与虚拟实验台；不是实际工具性能验收"
    document.save(target)
    print(target.relative_to(ROOT))


if __name__ == "__main__":
    markdown_document(FOLDER / "早上先看这里.md", FOLDER / "第二阶段平台交付与使用说明_20261009.docx")
    markdown_document(FOLDER / "模型结构与计算过程.md", FOLDER / "模型结构与计算过程_20261009.docx")
